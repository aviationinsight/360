#!/usr/bin/env python3
"""Verify that one valid appended row updates each dashboard by exactly one."""

from __future__ import annotations

import argparse
import json
import subprocess
import tempfile
from copy import copy
from pathlib import Path

from openpyxl import load_workbook

from amos_pipeline import validate as validate_amos
from aog_pipeline import SOURCE_SHEET as AOG_SHEET
from aog_pipeline import validate as validate_aog
from fak_pipeline import SOURCE_SHEET as FAK_SHEET
from fak_pipeline import validate as validate_fak
from snapshot_store import promote


CASES = {
    "AMOS": {
        "validator": validate_amos,
        "sheet": "Amos file",
        "total_key": "total_picking_lists",
        "unique_column": 6,
    },
    "AOG": {
        "validator": validate_aog,
        "sheet": AOG_SHEET,
        "total_key": "total_deliveries",
        "unique_column": 1,
    },
    "FAK": {
        "validator": validate_fak,
        "sheet": FAK_SHEET,
        "total_key": "total_movements",
        "unique_column": 1,
    },
}


def append_valid_row(source: Path, destination: Path, sheet_name: str, unique_column: int) -> None:
    workbook = load_workbook(source)
    sheet = workbook[sheet_name]
    source_row = sheet.max_row
    target_row = source_row + 1
    for column in range(1, sheet.max_column + 1):
        original = sheet.cell(source_row, column)
        target = sheet.cell(target_row, column)
        target.value = original.value
        if original.has_style:
            target._style = copy(original._style)
        target.number_format = original.number_format
    original_id = sheet.cell(source_row, unique_column).value
    sheet.cell(target_row, unique_column).value = f"{original_id}-AUTOMATION-TEST"
    raw_dir = destination.parent / "raw"
    raw_dir.mkdir(parents=True, exist_ok=True)
    raw_path = raw_dir / destination.name
    workbook.save(raw_path)
    subprocess.run(
        [
            "soffice", "--headless", "--convert-to", "xlsx",
            "--outdir", str(destination.parent), str(raw_path),
        ],
        check=True,
        capture_output=True,
        text=True,
    )
    if not destination.exists():
        raise RuntimeError("Spreadsheet recalculation did not produce the expected workbook.")


def run_case(dashboard: str, source: Path, root: Path, temp_dir: Path) -> dict:
    config = CASES[dashboard]
    validator = config["validator"]
    baseline = validator(source)
    baseline_total = baseline["kpis"][config["total_key"]]
    baseline_decision = promote(dashboard, baseline, root)

    appended_path = temp_dir / f"{dashboard.lower()}-valid-append.xlsx"
    append_valid_row(source, appended_path, config["sheet"], config["unique_column"])
    updated = validator(appended_path)
    updated_total = updated["kpis"][config["total_key"]]
    update_decision = promote(dashboard, updated, root)

    passed = (
        baseline_decision["accepted"]
        and update_decision["accepted"]
        and updated_total == baseline_total + 1
        and update_decision["current_checksum"] != baseline_decision["current_checksum"]
    )
    return {
        "dashboard": dashboard,
        "baseline_total": baseline_total,
        "updated_total": updated_total,
        "increase": updated_total - baseline_total,
        "updated_snapshot_promoted": update_decision["accepted"],
        "snapshot_changed": update_decision["current_checksum"] != baseline_decision["current_checksum"],
        "passed": passed,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--amos", required=True, type=Path)
    parser.add_argument("--aog", required=True, type=Path)
    parser.add_argument("--fak", required=True, type=Path)
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()

    with tempfile.TemporaryDirectory(prefix="sal-valid-append-") as temporary:
        temp_dir = Path(temporary)
        root = temp_dir / "snapshots"
        checks = [
            run_case("AMOS", args.amos, root, temp_dir),
            run_case("AOG", args.aog, root, temp_dir),
            run_case("FAK", args.fak, root, temp_dir),
        ]

    result = {"passed": all(item["passed"] for item in checks), "checks": checks}
    output = json.dumps(result, ensure_ascii=False, indent=2)
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(output + "\n", encoding="utf-8")
    print(output)
    raise SystemExit(0 if result["passed"] else 1)


if __name__ == "__main__":
    main()
