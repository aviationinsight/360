#!/usr/bin/env python3
"""Validate and normalize the SAL AOG Delivery workbook without changing it."""

from __future__ import annotations

import argparse
import json
import re
from collections import Counter
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


SOURCE_SHEET = "Automated "
EXPECTED_HEADERS = [
    "Picking List", "Issued At", "From Location", "To Location",
    "Time Issued", "Time Booked", "Received Time by Mechanics",
    "Driver Name/Mechanics", "On Time Status", "Minutes", "To Station",
    "From Store", "Year", "Month", "Day",
]


def clean(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def parse_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = clean(value).split(" ", 1)[0]
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def parse_time(value: Any) -> time | None:
    if isinstance(value, datetime):
        return value.time()
    if isinstance(value, time):
        return value
    if isinstance(value, timedelta):
        seconds = int(value.total_seconds()) % (24 * 60 * 60)
        return time(seconds // 3600, (seconds % 3600) // 60, seconds % 60)
    text = clean(value)
    for fmt in ("%H:%M:%S", "%H:%M"):
        try:
            return datetime.strptime(text, fmt).time()
        except ValueError:
            continue
    return None


def elapsed_minutes(booked: time, received: time) -> int:
    start = booked.hour * 60 + booked.minute + booked.second / 60
    end = received.hour * 60 + received.minute + received.second / 60
    if end < start:
        end += 24 * 60
    return round(end - start)


def canonical_station(value: Any) -> str:
    text = clean(value).upper()
    mapping = {
        "SHOP": "Shops", "SHOPS": "Shops",
        "SHIPPING": "Shipping", "TERMINAL": "Terminal",
        "LINE 601": "601", "601": "601",
    }
    return mapping.get(text, clean(value).title())


def validate(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    if SOURCE_SHEET not in workbook.sheetnames:
        return {"status": "rejected", "errors": [f"Required sheet {SOURCE_SHEET!r} was not found."]}

    sheet = workbook[SOURCE_SHEET]
    rows = sheet.iter_rows(values_only=True)
    headers = [clean(v) for v in next(rows)]
    if headers[:15] != EXPECTED_HEADERS:
        return {
            "status": "rejected",
            "errors": ["The AOG column structure does not match the approved template."],
            "expected_headers": EXPECTED_HEADERS,
            "received_headers": headers[:15],
        }

    accepted = rejected = 0
    rejection_reasons: Counter[str] = Counter()
    warnings: Counter[str] = Counter()
    dates: list[date] = []
    stations: Counter[str] = Counter()
    stores: Counter[str] = Counter()
    weekdays: Counter[str] = Counter()
    months: Counter[str] = Counter()
    source_locations: Counter[str] = Counter()
    seen_rows: set[tuple[str, ...]] = set()
    on_time = late = 0

    for row in rows:
        picking = clean(row[0])
        issued_at = parse_date(row[1])
        from_location = clean(row[2]).upper()
        to_location = clean(row[3])
        booked = parse_time(row[5])
        received = parse_time(row[6])
        source_status = clean(row[8]).lower()
        station = canonical_station(row[10])
        store = clean(row[11]).upper()

        reasons = []
        if not picking:
            reasons.append("missing picking list")
        if issued_at is None:
            reasons.append("invalid issued date")
        if not from_location or not to_location:
            reasons.append("missing route")
        if booked is None or received is None:
            reasons.append("invalid booked or received time")
        if not station or not store:
            reasons.append("missing operational dimension")

        calculated_status = ""
        minutes = None
        if booked is not None and received is not None:
            minutes = elapsed_minutes(booked, received)
            calculated_status = "on time" if minutes <= 30 else "late"
            if source_status not in {"on time", "late"}:
                reasons.append("invalid on-time status")
            elif source_status != calculated_status:
                reasons.append("on-time status does not match elapsed time")

        fingerprint = tuple(clean(value) for value in row[:15])
        if fingerprint in seen_rows:
            reasons.append("exact duplicate row")
        else:
            seen_rows.add(fingerprint)

        if reasons:
            rejected += 1
            rejection_reasons.update(reasons)
            continue

        accepted += 1
        dates.append(issued_at)
        stations[station] += 1
        stores[store] += 1
        source_locations[from_location] += 1
        months[issued_at.strftime("%b")] += 1
        weekdays[issued_at.strftime("%a")] += 1
        if calculated_status == "on time":
            on_time += 1
        else:
            late += 1
        if station != clean(row[10]):
            warnings["station label normalized"] += 1

    total = accepted
    return {
        "status": "accepted" if total and rejected == 0 else "accepted_with_exceptions" if total else "rejected",
        "source_file": path.name,
        "sheet": SOURCE_SHEET,
        "total_rows": accepted + rejected,
        "accepted_rows": accepted,
        "rejected_rows": rejected,
        "date_range": {
            "from": min(dates).isoformat() if dates else None,
            "to": max(dates).isoformat() if dates else None,
        },
        "kpis": {
            "total_deliveries": total,
            "on_time_deliveries": on_time,
            "late_deliveries": late,
            "on_time_percentage": round(on_time / total * 100, 2) if total else 0,
        },
        "breakdowns": {
            "destination_groups": stations,
            "from_stores": stores,
            "source_locations": source_locations,
            "months": months,
            "weekdays": weekdays,
        },
        "warnings": warnings,
        "rejection_reasons": rejection_reasons,
        "safe_to_replace_live_data": rejected == 0 and total > 0,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()
    report = validate(args.workbook)
    output = json.dumps(report, ensure_ascii=False, indent=2, default=dict)
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(output + "\n", encoding="utf-8")
    print(output)


if __name__ == "__main__":
    main()
