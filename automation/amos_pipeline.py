#!/usr/bin/env python3
"""Validate and normalize the SAL AMOS workbook without modifying the source file."""

from __future__ import annotations

import argparse
import json
import re
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


EXPECTED_HEADERS = [
    "YearID", "MonthID", "Month Name", "Day ID", "Shift Time",
    "Picking List", "Prio", "Station", "Store", "Issue Date",
    "Booked Date", "Booked Date", "Mech Sign", "Booking Sign",
    "Time Taken in min", "Flag", "Remarks", "Discription",
    "Shift Name", "Time Slot",
]
VALID_PRIORITIES = {"AOG", "Normal"}
VALID_FLAGS = {
    "AOG on time", "AOG exceeding 30",
    "Normal on time", "Normal exceeding 60",
}


def clean(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def parse_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = clean(value).split(" ", 1)[0]
    for fmt in ("%d-%m-%Y", "%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def parse_component_date(year_value: Any, month_value: Any, day_value: Any) -> date | None:
    try:
        return date(
            int(float(clean(year_value))),
            int(float(clean(month_value))),
            int(float(clean(day_value))),
        )
    except (TypeError, ValueError):
        return None


def parse_minutes(value: Any) -> float | None:
    try:
        return float(clean(value).replace(",", ""))
    except ValueError:
        return None


def validate(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    if "Amos file" not in workbook.sheetnames:
        return {"status": "rejected", "errors": ["Required sheet 'Amos file' was not found."]}

    sheet = workbook["Amos file"]
    rows = sheet.iter_rows(values_only=True)
    headers = [clean(v) for v in next(rows)]
    if headers[:20] != EXPECTED_HEADERS:
        return {
            "status": "rejected",
            "errors": ["The AMOS column structure does not match the approved template."],
            "expected_headers": EXPECTED_HEADERS,
            "received_headers": headers[:20],
        }

    accepted = 0
    rejected = 0
    warnings: Counter[str] = Counter()
    rejection_reasons: Counter[str] = Counter()
    priorities: Counter[str] = Counter()
    flags: Counter[str] = Counter()
    stations: Counter[str] = Counter()
    stores: Counter[str] = Counter()
    shifts: Counter[str] = Counter()
    dates: list[date] = []
    seen_rows: set[tuple[str, ...]] = set()
    picking_counts: Counter[str] = Counter()
    aog_total = aog_on_time = normal_total = normal_on_time = 0

    for row_number, row in enumerate(rows, start=2):
        picking = clean(row[5])
        priority = clean(row[6])
        station = clean(row[7]).upper()
        store = clean(row[8]).upper()
        operation_date = parse_component_date(row[0], row[1], row[3])
        minutes = parse_minutes(row[14])
        flag = clean(row[15])
        shift = clean(row[18])

        reasons = []
        if not picking:
            reasons.append("missing picking list")
        if operation_date is None:
            reasons.append("invalid YearID, MonthID, or Day ID")
        if priority not in VALID_PRIORITIES:
            reasons.append("invalid priority")
        if flag not in VALID_FLAGS:
            reasons.append("invalid performance flag")
        if minutes is None or minutes < 0:
            reasons.append("invalid elapsed minutes")
        if not station or not store or not shift:
            reasons.append("missing operational dimension")

        if reasons:
            rejected += 1
            rejection_reasons.update(reasons)
            continue

        fingerprint = tuple(clean(value) for value in row[:20])
        if fingerprint in seen_rows:
            warnings["exact duplicate row retained for review"] += 1
        seen_rows.add(fingerprint)

        accepted += 1
        picking_counts[picking] += 1
        dates.append(operation_date)
        priorities[priority] += 1
        flags[flag] += 1
        stations[station] += 1
        stores[store] += 1
        shifts[shift] += 1

        if not flag.startswith(priority):
            warnings["source priority differs from final reviewed flag"] += 1

        if flag.startswith("AOG"):
            aog_total += 1
            aog_on_time += int(flag == "AOG on time")
        else:
            normal_total += 1
            normal_on_time += int(flag == "Normal on time")

    status = "accepted" if accepted and rejected == 0 else "accepted_with_exceptions" if accepted else "rejected"
    return {
        "status": status,
        "source_file": path.name,
        "sheet": "Amos file",
        "total_rows": accepted + rejected,
        "accepted_rows": accepted,
        "rejected_rows": rejected,
        "date_range": {
            "from": min(dates).isoformat() if dates else None,
            "to": max(dates).isoformat() if dates else None,
        },
        "kpis": {
            "total_picking_lists": accepted,
            "aog_total": aog_total,
            "aog_on_time": aog_on_time,
            "aog_on_time_percentage": round(aog_on_time / aog_total * 100, 2) if aog_total else 0,
            "normal_total": normal_total,
            "normal_on_time": normal_on_time,
            "normal_on_time_percentage": round(normal_on_time / normal_total * 100, 2) if normal_total else 0,
        },
        "breakdowns": {
            "priorities": priorities,
            "flags": flags,
            "stations": stations,
            "top_stores": dict(stores.most_common(15)),
            "shifts": shifts,
        },
        "warnings": {
            **warnings,
            "picking list IDs appearing more than once": sum(value > 1 for value in picking_counts.values()),
            "additional occurrences of repeated picking list IDs": sum(value - 1 for value in picking_counts.values() if value > 1),
        },
        "rejection_reasons": rejection_reasons,
        "safe_to_replace_live_data": rejected == 0 and accepted > 0,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()
    report = validate(args.workbook)
    output = json.dumps(report, ensure_ascii=False, indent=2, default=dict)
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(output + "\n", encoding="utf-8")
    print(output)


if __name__ == "__main__":
    main()
