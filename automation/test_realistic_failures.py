#!/usr/bin/env python3
"""Exercise realistic workbook failures against the last-known-good gate."""

from __future__ import annotations

import argparse
import json
import tempfile
from pathlib import Path

from openpyxl import load_workbook

from amos_pipeline import validate as validate_amos
from aog_pipeline import SOURCE_SHEET as AOG_SHEET
from aog_pipeline import validate as validate_aog
from fak_pipeline import SOURCE_SHEET as FAK_SHEET
from fak_pipeline import validate as validate_fak
from snapshot_store import checksum, promote


CASES = {
    "AMOS": {"validator": validate_amos, "sheet": "Amos file", "date_cell": "A2"},
    "AOG": {"validator": validate_aog, "sheet": AOG_SHEET, "date_cell": "B2"},
    "FAK": {"validator": validate_fak, "sheet": FAK_SHEET, "date_cell": "B2"},
}


def save_mutation(source: Path, destination: Path, sheet_name: str, mutation: str, date_cell: str) -> None:
    workbook = load_workbook(source)
    sheet = workbook[sheet_name]
    if mutation == "missing_header":
        sheet.cell(1, 1).value = "REMOVED_REQUIRED_HEADER"
    elif mutation == "invalid_date":
        sheet[date_cell] = "not-a-date"
    else:
        raise ValueError(f"Unknown mutation: {mutation}")
    workbook.save(destination)


def run_case(dashboard: str, source: Path, root: Path, temp_dir: Path) -> dict:
    config = CASES[dashboard]
    validator = config["validator"]
    valid_report = validator(source)
    valid_decision = promote(dashboard, valid_report, root)
    current_path = root / "current" / f"{dashboard.lower()}.json"
    baseline_hash = checksum(current_path)

    mutations = []
    for mutation in ("missing_header", "invalid_date"):
        broken_path = temp_dir / f"{dashboard.lower()}-{mutation}.xlsx"
        save_mutation(source, broken_path, config["sheet"], mutation, config["date_cell"])
        report = validator(broken_path)
        decision = promote(dashboard, report, root)
        current_hash = checksum(current_path)
        rejected = not decision["accepted"]
        preserved = current_hash == baseline_hash
        mutations.append({
            "mutation": mutation,
            "validation_status": report.get("status"),
            "rejected": rejected,
            "last_good_snapshot_unchanged": preserved,
            "passed": rejected and preserved,
        })

    passed = valid_decision["accepted"] and all(item["passed"] for item in mutations)
    return {
        "dashboard": dashboard,
        "valid_baseline_saved": valid_decision["accepted"],
        "mutations": mutations,
        "passed": passed,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--amos", required=True, type=Path)
    parser.add_argument("--aog", required=True, type=Path)
    parser.add_argument("--fak", required=True, type=Path)
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()

    with tempfile.TemporaryDirectory(prefix="sal-realistic-failures-") as temporary:
        temp_dir = Path(temporary)
        root = temp_dir / "snapshots"
        checks = [
            run_case("AMOS", args.amos, root, temp_dir),
            run_case("AOG", args.aog, root, temp_dir),
            run_case("FAK", args.fak, root, temp_dir),
        ]

    result = {"passed": all(item["passed"] for item in checks), "checks": checks}
    output = json.dumps(result, ensure_ascii=False, indent=2)
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(output + "\n", encoding="utf-8")
    print(output)
    raise SystemExit(0 if result["passed"] else 1)


if __name__ == "__main__":
    main()
