#!/usr/bin/env python3
"""Validate and normalize the SAL FAK workbook without changing the source."""

from __future__ import annotations

import argparse
import json
import re
from collections import Counter
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


SOURCE_SHEET = "FAK AutomatedNew "
EXPECTED_HEADERS = [
    "SN", "DATE", "FLIGHT NO", "ACFT REGISTRATION", "CONTAINER",
    "TYPE", "TIME", "SIGN TIME", "STATUS", "SHIFT", "REMARK",
    "DESTINATION", "Aircraft Type",
]
VALID_MOVEMENTS = {"ARRIVAL", "DEPARTURE", "MAINTENANCE"}
VALID_STATUSES = {"ON TIME", "DELAY", "MISSED"}
VALID_SHIFTS = {"DAY", "TWI", "GRV"}


def clean(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def parse_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = clean(value).split(" ", 1)[0]
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def valid_time(value: Any, allow_maintenance: bool) -> bool:
    if isinstance(value, (datetime, time, timedelta)):
        return True
    text = clean(value).upper()
    if allow_maintenance and text in {"-", "MAINTENANCE"}:
        return True
    if text == "-":
        return True
    for fmt in ("%H:%M:%S", "%H:%M"):
        try:
            datetime.strptime(text, fmt)
            return True
        except ValueError:
            continue
    return False


def canonical_flight(value: Any) -> str:
    text = clean(value).upper()
    return re.sub(r"^SV0+(\d+)$", r"SV\1", text)


def validate(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    if SOURCE_SHEET not in workbook.sheetnames:
        return {"status": "rejected", "errors": [f"Required sheet {SOURCE_SHEET!r} was not found."]}

    sheet = workbook[SOURCE_SHEET]
    rows = sheet.iter_rows(values_only=True)
    headers = [clean(v) for v in next(rows)]
    if headers[:13] != EXPECTED_HEADERS:
        return {
            "status": "rejected",
            "errors": ["The FAK column structure does not match the approved template."],
            "expected_headers": EXPECTED_HEADERS,
            "received_headers": headers[:13],
        }

    accepted = rejected = 0
    rejection_reasons: Counter[str] = Counter()
    warnings: Counter[str] = Counter()
    dates: list[date] = []
    movements: Counter[str] = Counter()
    statuses: Counter[str] = Counter()
    shifts: Counter[str] = Counter()
    destinations: Counter[str] = Counter()
    aircraft_types: Counter[str] = Counter()
    flights: Counter[str] = Counter()
    seen_rows: set[tuple[str, ...]] = set()

    for row in rows:
        serial = clean(row[0])
        operation_date = parse_date(row[1])
        raw_flight = clean(row[2]).upper()
        flight = canonical_flight(row[2])
        registration = clean(row[3]).upper()
        container = clean(row[4]).upper()
        raw_movement = clean(row[5])
        movement = raw_movement.upper()
        raw_status = clean(row[8])
        status = raw_status.upper()
        raw_shift = clean(row[9])
        shift = raw_shift.upper()
        destination = clean(row[11]).upper()
        aircraft_type = clean(row[12]).upper()
        expected_aircraft_type = "CARGO" if destination == "CARGO" else "COMMERCIAL"

        reasons = []
        if operation_date is None:
            reasons.append("invalid operation date")
        if not flight or not registration or not container:
            reasons.append("missing flight identity")
        if movement not in VALID_MOVEMENTS:
            reasons.append("invalid movement type")
        if status not in VALID_STATUSES:
            reasons.append("invalid status")
        if shift not in VALID_SHIFTS:
            reasons.append("invalid shift")
        if not destination:
            reasons.append("missing destination")
        if aircraft_type != expected_aircraft_type:
            reasons.append("aircraft type does not match destination rule")
        if not valid_time(row[6], movement == "MAINTENANCE"):
            reasons.append("invalid operation time")

        fingerprint = tuple(clean(value) for value in row[:13])
        if fingerprint in seen_rows:
            reasons.append("exact duplicate row")
        else:
            seen_rows.add(fingerprint)

        if reasons:
            rejected += 1
            rejection_reasons.update(reasons)
            continue

        accepted += 1
        dates.append(operation_date)
        movements[movement] += 1
        statuses[status] += 1
        shifts[shift] += 1
        destinations[destination] += 1
        aircraft_types[aircraft_type] += 1
        flights[flight] += 1
        if not serial:
            warnings["missing source serial; internal ID will be generated"] += 1
        if raw_movement != movement or raw_status != status or raw_shift != shift:
            warnings["label whitespace or casing normalized"] += 1
        if raw_flight != flight:
            warnings["flight number leading zero normalized"] += 1

    total = accepted
    on_time = statuses["ON TIME"]
    return {
        "status": "accepted" if total and rejected == 0 else "accepted_with_exceptions" if total else "rejected",
        "source_file": path.name,
        "sheet": SOURCE_SHEET,
        "total_rows": accepted + rejected,
        "accepted_rows": accepted,
        "rejected_rows": rejected,
        "date_range": {
            "from": min(dates).isoformat() if dates else None,
            "to": max(dates).isoformat() if dates else None,
        },
        "kpis": {
            "total_movements": total,
            "passenger_aircraft": aircraft_types["COMMERCIAL"],
            "cargo_aircraft": aircraft_types["CARGO"],
            "on_time_movements": on_time,
            "delayed_movements": statuses["DELAY"],
            "missed_flights": statuses["MISSED"],
            "on_time_percentage": round(on_time / total * 100, 2) if total else 0,
        },
        "breakdowns": {
            "movement_types": movements,
            "statuses": statuses,
            "shifts": shifts,
            "top_destinations": dict(destinations.most_common(15)),
            "top_flights": dict(flights.most_common(15)),
        },
        "warnings": warnings,
        "rejection_reasons": rejection_reasons,
        "safe_to_replace_live_data": rejected == 0 and total > 0,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()
    report = validate(args.workbook)
    output = json.dumps(report, ensure_ascii=False, indent=2, default=dict)
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(output + "\n", encoding="utf-8")
    print(output)


if __name__ == "__main__":
    main()
